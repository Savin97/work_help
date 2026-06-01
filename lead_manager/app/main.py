import io
import csv
import base64
from datetime import date
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, Depends, Request, Form, UploadFile, File, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session

from .database import engine, get_db
from . import models, crud, schemas

models.Base.metadata.create_all(bind=engine)

app = FastAPI(title="Lead Manager")

TEMPLATES_DIR = Path(__file__).parent / "templates"
templates = Jinja2Templates(directory=str(TEMPLATES_DIR))

INDUSTRIES = ["pharma", "university", "hospital", "food", "environmental", "government", "other"]
STATUSES = ["new", "contacted", "engaged", "proposal_sent", "won", "lost", "dormant"]
CONTACT_TYPES = ["call", "email", "meeting", "linkedin", "other"]

INDUSTRY_LABELS = {
    "pharma": "Pharma / Biotech",
    "university": "University / Research",
    "hospital": "Hospital / Clinical",
    "food": "Food & Beverage",
    "environmental": "Environmental",
    "government": "Government / Defense",
    "other": "Other",
}
STATUS_LABELS = {
    "new": "New",
    "contacted": "Contacted",
    "engaged": "Engaged",
    "proposal_sent": "Proposal Sent",
    "won": "Won",
    "lost": "Lost",
    "dormant": "Dormant",
}


# ── Dashboard ────────────────────────────────────────────────────────────────

@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request, db: Session = Depends(get_db)):
    stats = crud.get_dashboard_stats(db)
    return templates.TemplateResponse(request, "dashboard.html", {
        "stats": stats,
        "status_labels": STATUS_LABELS,
        "industry_labels": INDUSTRY_LABELS,
        "today": date.today(),
    })


# ── Leads list ───────────────────────────────────────────────────────────────

@app.get("/leads", response_class=HTMLResponse)
def leads_list(
    request: Request,
    status: Optional[str] = None,
    industry: Optional[str] = None,
    search: Optional[str] = None,
    msg: Optional[str] = None,
    db: Session = Depends(get_db),
):
    leads = crud.get_leads(db, status=status, industry=industry, search=search)
    return templates.TemplateResponse(request, "leads.html", {
        "leads": leads,
        "industries": INDUSTRIES,
        "statuses": STATUSES,
        "industry_labels": INDUSTRY_LABELS,
        "status_labels": STATUS_LABELS,
        "filter_status": status,
        "filter_industry": industry,
        "search": search or "",
        "today": date.today().isoformat(),
        "msg": msg,
    })


# ── Lead create ───────────────────────────────────────────────────────────────

@app.get("/leads/new", response_class=HTMLResponse)
def lead_new_form(request: Request):
    return templates.TemplateResponse(request, "lead_form.html", {
        "lead": None,
        "industries": INDUSTRIES,
        "statuses": STATUSES,
        "industry_labels": INDUSTRY_LABELS,
        "status_labels": STATUS_LABELS,
        "action": "/leads/new",
        "title": "Add Lead",
    })


@app.post("/leads/new")
def lead_create(
    request: Request,
    company: str = Form(...),
    contact_name: str = Form(""),
    phone: str = Form(""),
    email: str = Form(""),
    industry: str = Form("other"),
    equipment_interest: str = Form(""),
    status: str = Form("new"),
    source: str = Form(""),
    notes: str = Form(""),
    next_followup_date: str = Form(""),
    db: Session = Depends(get_db),
):
    lead_data = schemas.LeadCreate(
        company=company,
        contact_name=contact_name or None,
        phone=phone or None,
        email=email or None,
        industry=industry or None,
        equipment_interest=equipment_interest or None,
        status=status,
        source=source or None,
        notes=notes or None,
        next_followup_date=date.fromisoformat(next_followup_date) if next_followup_date else None,
    )
    lead = crud.create_lead(db, lead_data)
    return RedirectResponse(f"/leads/{lead.id}?msg=created", status_code=303)


# ── Lead detail & edit ────────────────────────────────────────────────────────

@app.get("/leads/{lead_id}", response_class=HTMLResponse)
def lead_detail(
    request: Request,
    lead_id: int,
    msg: Optional[str] = None,
    db: Session = Depends(get_db),
):
    lead = crud.get_lead(db, lead_id)
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    return templates.TemplateResponse(request, "lead_detail.html", {
        "lead": lead,
        "industries": INDUSTRIES,
        "statuses": STATUSES,
        "contact_types": CONTACT_TYPES,
        "industry_labels": INDUSTRY_LABELS,
        "status_labels": STATUS_LABELS,
        "today": date.today().isoformat(),
        "msg": msg,
    })


@app.post("/leads/{lead_id}/edit")
def lead_edit(
    lead_id: int,
    company: str = Form(...),
    contact_name: str = Form(""),
    phone: str = Form(""),
    email: str = Form(""),
    industry: str = Form("other"),
    equipment_interest: str = Form(""),
    status: str = Form("new"),
    source: str = Form(""),
    notes: str = Form(""),
    next_followup_date: str = Form(""),
    db: Session = Depends(get_db),
):
    lead_data = schemas.LeadUpdate(
        company=company,
        contact_name=contact_name or None,
        phone=phone or None,
        email=email or None,
        industry=industry or None,
        equipment_interest=equipment_interest or None,
        status=status,
        source=source or None,
        notes=notes or None,
        next_followup_date=date.fromisoformat(next_followup_date) if next_followup_date else None,
    )
    crud.update_lead(db, lead_id, lead_data)
    return RedirectResponse(f"/leads/{lead_id}?msg=saved", status_code=303)


@app.post("/leads/{lead_id}/delete")
def lead_delete(lead_id: int, db: Session = Depends(get_db)):
    crud.delete_lead(db, lead_id)
    return RedirectResponse("/leads?msg=deleted", status_code=303)


# ── Contact log ───────────────────────────────────────────────────────────────

@app.post("/leads/{lead_id}/log")
def add_log(
    lead_id: int,
    log_date: str = Form(...),
    contact_type: str = Form(...),
    summary: str = Form(""),
    outcome: str = Form(""),
    next_followup_date: str = Form(""),
    db: Session = Depends(get_db),
):
    log_data = schemas.ContactLogCreate(
        date=date.fromisoformat(log_date),
        contact_type=contact_type,
        summary=summary or None,
        outcome=outcome or None,
        next_followup_date=date.fromisoformat(next_followup_date) if next_followup_date else None,
    )
    crud.add_contact_log(db, lead_id, log_data)
    return RedirectResponse(f"/leads/{lead_id}?msg=logged", status_code=303)


@app.post("/logs/{log_id}/delete")
def delete_log(log_id: int, lead_id: int = Form(...), db: Session = Depends(get_db)):
    crud.delete_contact_log(db, log_id)
    return RedirectResponse(f"/leads/{lead_id}?msg=log_deleted", status_code=303)


# ── Import ────────────────────────────────────────────────────────────────────

@app.get("/import", response_class=HTMLResponse)
def import_page(request: Request, msg: Optional[str] = None, created: Optional[int] = None, skipped: Optional[int] = None):
    return templates.TemplateResponse(request, "import.html", {
        "msg": msg,
        "created": created,
        "skipped": skipped,
        "industries": INDUSTRIES,
        "industry_labels": INDUSTRY_LABELS,
    })


@app.post("/import/preview", response_class=HTMLResponse)
async def import_preview(
    request: Request,
    file: UploadFile = File(...),
):
    content = await file.read()
    filename = file.filename or ""
    rows = []
    headers = []

    try:
        if filename.endswith(".xlsx"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content))
            ws = wb.active
            all_rows = list(ws.iter_rows(values_only=True))
            if all_rows:
                headers = [str(h) if h is not None else "" for h in all_rows[0]]
                rows = [list(r) for r in all_rows[1:51]]
        else:
            text = content.decode("utf-8-sig", errors="replace")
            reader = csv.reader(io.StringIO(text))
            all_rows = list(reader)
            if all_rows:
                headers = all_rows[0]
                rows = all_rows[1:51]
    except Exception as e:
        return templates.TemplateResponse(request, "import.html", {
            "msg": f"Error reading file: {e}",
            "industries": INDUSTRIES,
            "industry_labels": INDUSTRY_LABELS,
        })

    return templates.TemplateResponse(request, "import_preview.html", {
        "headers": headers,
        "rows": rows,
        "filename": filename,
        "file_content": base64.b64encode(content).decode("ascii"),
        "industries": INDUSTRIES,
        "industry_labels": INDUSTRY_LABELS,
        "fields": ["company", "contact_name", "phone", "email", "industry", "equipment_interest", "source", "notes", "next_followup_date"],
    })


@app.post("/import/do")
async def import_do(
    request: Request,
    db: Session = Depends(get_db),
):
    form = await request.form()
    file_content_b64 = form.get("file_content", "")
    filename = form.get("filename", "")
    default_status = form.get("default_status", "dormant")
    default_industry = form.get("default_industry", "other")

    mapping = {}
    for key, val in form.items():
        if key.startswith("map_") and val and val != "__skip__":
            field = key[4:]
            mapping[field] = int(val)

    rows = []
    try:
        content = base64.b64decode(file_content_b64)
        if filename.endswith(".xlsx"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content))
            ws = wb.active
            all_rows = list(ws.iter_rows(values_only=True))
            rows = [list(r) for r in all_rows[1:]]
        else:
            text = content.decode("utf-8-sig", errors="replace")
            reader = csv.reader(io.StringIO(text))
            all_rows = list(reader)
            rows = all_rows[1:]
    except Exception as e:
        return RedirectResponse("/import?msg=parse_error", status_code=303)

    leads_to_create = []
    for row in rows:
        def get_val(field, _row=row):
            idx = mapping.get(field)
            if idx is None or idx >= len(_row):
                return None
            v = _row[idx]
            return str(v).strip() if v is not None else None

        company = get_val("company")
        if not company:
            continue
        nf = get_val("next_followup_date")
        nf_date = None
        if nf:
            from datetime import datetime
            for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
                try:
                    nf_date = datetime.strptime(nf, fmt).date()
                    break
                except ValueError:
                    pass

        industry = get_val("industry") or default_industry
        if industry not in INDUSTRIES:
            industry = default_industry

        leads_to_create.append(schemas.LeadCreate(
            company=company,
            contact_name=get_val("contact_name"),
            phone=get_val("phone"),
            email=get_val("email"),
            industry=industry,
            equipment_interest=get_val("equipment_interest"),
            status=default_status,
            source=get_val("source"),
            notes=get_val("notes"),
            next_followup_date=nf_date,
        ))

    created, skipped = crud.bulk_create_leads(db, leads_to_create)
    return RedirectResponse(f"/import?msg=done&created={created}&skipped={skipped}", status_code=303)


# ── Templates page ────────────────────────────────────────────────────────────

@app.get("/templates", response_class=HTMLResponse)
def outreach_templates(request: Request):
    return templates.TemplateResponse(request, "templates.html", {})


# ── Discovery guide ───────────────────────────────────────────────────────────

@app.get("/discovery", response_class=HTMLResponse)
def discovery(request: Request):
    return templates.TemplateResponse(request, "discovery.html", {})
